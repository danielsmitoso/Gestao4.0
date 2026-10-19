"""Teste automatizado (AppTest, já incluso no Streamlit -- sem dependência nova).

Como rodar (na pasta do projeto):
    python tests/test_app.py
(também funciona com `pytest tests/`, se o pytest estiver instalado)

Os testes usam bancos SQLite temporários: nunca tocam no banco do time.
"""
import io
import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

import openpyxl  # noqa: E402
from pptx import Presentation  # noqa: E402
from streamlit.testing.v1 import AppTest  # noqa: E402

from modules import banco as bd  # noqa: E402
from modules import exemplo, export_excel, export_pptx, qualidade, regras  # noqa: E402
from modules import historico as hist  # noqa: E402
from modules import importador as imp  # noqa: E402

TIMEOUT = 120
HOJE = date(2026, 10, 19)
PLANILHA_REAL = RAIZ / "Gestão 4.0 - NTec.xlsx"


def _exemplo():
    res = imp.ler_planilha(exemplo.gerar_planilha_exemplo(HOJE))
    return res, regras.classificar(res.acoes, HOJE)


def _planilha(versao: int):
    return imp.ler_planilha(exemplo.gerar_planilha_exemplo(HOJE, versao)).acoes


def _gravar(b: bd.Banco, versao: int, posicao: date, recusar=frozenset()):
    plano = hist.planejar(_planilha(versao), b.carregar(), set(recusar))
    b.gravar_importacao(plano, posicao, f"v{versao}.xlsx", "Teste")
    return plano


def _banco_duas_semanas() -> bd.Banco:
    b = bd.conectar("sqlite://")
    _gravar(b, 1, HOJE - timedelta(days=7))
    _gravar(b, 2, HOJE)
    return b


# ---------------------------------------------------------------------------
# Leitura e regras
# ---------------------------------------------------------------------------
def test_importador_le_formato_atual():
    res, df = _exemplo()
    assert res.abas_lidas == ["Fábrica de Placas", "Centralização", "Novos Negócios"], res.abas_lidas
    assert res.abas_vazias == ["Modelo vazio"], "aba só com 'AÇÕES' deve ser ignorada"
    assert "Instruções" in res.abas_ignoradas
    assert res.avisos == [], "aba só com o modelo não deve gerar aviso de estrutura"
    assert len(df) == 21, len(df)
    assert df["entrega"].nunique() == 6
    sub = df[df["acao"] == "Catalogar placas prontas"].iloc[0]
    assert (sub["entrega"], sub["subgrupo"]) == ("Plano de produção", "Estoque")
    assert not df["acao"].isin(["AÇÕES", "Estoque", "Execução"]).any(), "subgrupo/modelo não é ação"
    assert df.loc[df["acao"] == "Integração do novo supervisor", "comentario"].item() == "Agenda de integração enviada"


def test_sinais_por_prazo():
    _, df = _exemplo()
    sinal = dict(zip(df["acao"], df["sinal"]))
    assert sinal["Troca de lâmpadas do galpão"] == regras.ATRASADA        # -3 dias
    assert sinal["Formulário Hub Norte"] == regras.ATRASADA               # -1 dia
    assert sinal["Ar-condicionado da produção"] == regras.VENCE           # +2
    assert sinal["Plano de impressão semanal"] == regras.VENCE            # +6
    assert sinal["Planejamento de entregas às unidades"] == regras.NO_PRAZO  # +15
    assert sinal["Revisão dos compressores"] == regras.SEM_PRAZO
    assert sinal["Catalogar placas prontas"] == regras.CONCLUIDA          # concluída, mesmo com data passada
    df2 = regras.classificar(df.drop(columns=["sinal"]), date(2026, 11, 30))
    assert dict(zip(df2["acao"], df2["sinal"]))["Planejamento de entregas às unidades"] == regras.ATRASADA


def test_avanco_ponderado_igual_planilha():
    _, df = _exemplo()
    g = df[df["entrega"] == "Equipe"]  # pesos 2,3,1 com 100%, 50%, 0%
    assert abs(regras.avanco(g) - (2 * 1 + 3 * 0.5 + 1 * 0) / 6) < 1e-9
    g = df[df["entrega"] == "Plano de produção"]  # ação sem peso não entra (SUMPRODUCT/SUM da planilha)
    com_peso = g[g["peso"].notna()]
    esperado = (com_peso["peso"] * com_peso["pct"]).sum() / com_peso["peso"].sum()
    assert abs(regras.avanco(g) - esperado) < 1e-9


def test_farol_das_entregas():
    _, df = _exemplo()
    ent = regras.resumo_entregas(df).set_index("entrega")["farol"]
    assert ent["Manutenção da fábrica"] == "vermelho"     # tem atrasada
    assert ent["Equipe"] == "amarelo"                     # vence em 4 dias, nada atrasado
    assert ent["Projeto Alfa - Orçamento"] == "amarelo"   # 2 de 3 sem prazo
    concluida = df[df["entrega"] == "Equipe"].assign(sinal=regras.CONCLUIDA, aberta=False)
    assert regras.farol(concluida) == "concluida"


def test_responsaveis_e_qualidade():
    _, df = _exemplo()
    carga = regras.carga_responsaveis(df).set_index("responsavel")
    assert carga.loc["Bruno", "acoes"] == 5, "ação 'Bruno, Carla' e 'Eva / Bruno' contam para Bruno"
    assert bool(carga.loc["Facilities", "externo"])
    nomes = {c["checagem"]: c["qtd"] for c in qualidade.checar(df)}
    assert nomes["Ação em aberto sem responsável"] == 2
    assert nomes["Mais de um responsável na mesma ação"] == 2
    assert nomes["Área de apoio como responsável"] == 2


def test_exportacoes_abrem():
    _, df = _exemplo()
    prs = Presentation(io.BytesIO(export_pptx.gerar_relatorio(df, "Teste", HOJE)))
    assert len(prs.slides) >= 6
    wb = openpyxl.load_workbook(io.BytesIO(export_excel.gerar_base(df, HOJE)))
    assert wb.sheetnames[:3] == ["Ações", "Entregas", "Qualidade dos dados"]
    assert wb["Ações"].max_row == len(df) + 1
    com_hist = regras.classificar(hist.estado_atual(_banco_duas_semanas().carregar()), HOJE)
    Presentation(io.BytesIO(export_pptx.gerar_relatorio(com_hist, "Teste", HOJE)))
    wb = openpyxl.load_workbook(io.BytesIO(export_excel.gerar_base(com_hist, HOJE)))
    assert "Prazo adiado (vezes)" in [c.value for c in wb["Ações"][1]]


def test_planilha_real_se_existir():
    """Roda só na máquina que tem a planilha real (ela não vai para o repositório)."""
    if not PLANILHA_REAL.exists():
        print("   (planilha real ausente -- pulado)")
        return
    res = imp.ler_planilha(PLANILHA_REAL)
    df = regras.classificar(res.acoes, HOJE)
    assert len(df) > 0 and "Fábrica de Placas" in res.abas_lidas
    assert set(df["status"]) <= set(regras.ORDEM_STATUS)
    # grava duas vezes num banco temporário: a 2ª não pode mudar nada
    b = bd.conectar("sqlite://")
    b.gravar_importacao(hist.planejar(res.acoes, b.carregar()), HOJE, "real.xlsx", "Teste")
    r = hist.planejar(res.acoes, b.carregar()).resumo()
    assert r["sem_mudanca"] == len(df) and r["novas"] == 0, r
    export_pptx.gerar_relatorio(df, "Real", HOJE)
    export_excel.gerar_base(df, HOJE)


# ---------------------------------------------------------------------------
# Histórico (memória)
# ---------------------------------------------------------------------------
def test_historico_compara_importacoes():
    b = bd.conectar("sqlite://")
    p1 = _gravar(b, 1, HOJE - timedelta(days=7))
    assert p1.resumo()["novas"] == 21
    p2 = hist.planejar(_planilha(2), b.carregar())
    r = p2.resumo()
    assert (r["novas"], r["atualizadas"], r["sairam"], r["renomeadas"]) == (1, 5, 1, 1), r
    assert p2.renomeacoes[0]["de"].endswith("Catalogar placas prontas")
    mud = dict(zip(p2.linhas["acao"], p2.linhas["mudancas"]))
    assert "Prazo: 21/10/2026 → 28/10/2026" in mud["Ar-condicionado da produção"]
    b.gravar_importacao(p2, HOJE, "v2.xlsx", "Teste")
    r = hist.planejar(_planilha(2), b.carregar()).resumo()  # mesma planilha de novo = nada muda
    assert r["sem_mudanca"] == r["acoes"] and r["sairam"] == 0, r


def test_estado_atual_e_eventos():
    base = _banco_duas_semanas().carregar()
    atual = hist.estado_atual(base).set_index("acao")
    assert len(atual) == 22, "21 da v1 + 1 nova; a renomeada continua sendo a mesma demanda"
    assert bool(atual.loc["Integração do novo supervisor", "concluida_na_ultima"])
    assert atual.loc["Integração do novo supervisor", "concluida_em"] == HOJE
    assert not bool(atual.loc["Catalogar as placas prontas", "concluida_na_ultima"]), "já estava concluída"
    assert atual.loc["Ar-condicionado da produção", "vezes_prazo_alterado"] == 1
    assert atual.loc["Ar-condicionado da produção", "prazo_original"] == HOJE + timedelta(days=2)
    assert not bool(atual.loc["Escala de férias do 1º semestre", "ativa"])
    assert atual.loc["Manutenção de EPS", "incluida_em"] == HOJE
    tipos = set(hist.eventos(base)["tipo"])
    assert {"Renomeada", "Status", "Prazo", "Nova demanda", "Saiu da planilha"} <= tipos, tipos
    assert list(hist.evolucao(base)["data_posicao"]) == [HOJE - timedelta(days=7), HOJE]
    pauta = regras.pauta(regras.classificar(hist.estado_atual(base), HOJE), HOJE)
    assert set(pauta["concluido"]["acao"]) == {"Integração do novo supervisor", "Formulário Hub Norte"}


def test_recusar_renomeacao_e_voltar():
    b = bd.conectar("sqlite://")
    _gravar(b, 1, HOJE - timedelta(days=14))
    recusado = hist.planejar(_planilha(2), b.carregar()).renomeacoes[0]["demanda_id"]
    p = _gravar(b, 2, HOJE - timedelta(days=7), {recusado})
    assert (p.resumo()["novas"], p.resumo()["sairam"]) == (2, 2), p.resumo()
    # planilha antiga de novo: as que saíram "voltam" (mesma demanda, histórico preservado)
    assert hist.planejar(_planilha(1), b.carregar()).contagem("voltou") == 2


def test_desfazer_ultima_importacao():
    b = _banco_duas_semanas()
    assert b.desfazer_ultima("Teste") is not None
    base = b.carregar()
    assert len(base["importacoes"]) == 1
    atual = hist.estado_atual(base).set_index("acao")
    assert atual.loc["Integração do novo supervisor", "status"] == "Em andamento"
    assert "Manutenção de EPS" not in atual.index
    assert bool(atual["ativa"].all())


def test_banco_em_arquivo_persiste():
    with tempfile.TemporaryDirectory() as tmp:
        url = f"sqlite:///{Path(tmp, 'h.db').as_posix()}"
        b = bd.conectar(url)
        _gravar(b, 1, HOJE)
        b.engine.dispose()
        b2 = bd.conectar(url)
        assert len(b2.carregar()["importacoes"]) == 1
        b2.engine.dispose()


# ---------------------------------------------------------------------------
# App
# ---------------------------------------------------------------------------
def _app(secrets: dict) -> AppTest:
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=TIMEOUT)
    for k, v in secrets.items():
        at.secrets[k] = v
    at.run()
    assert not at.exception, at.exception
    return at


def _metricas(at) -> dict:
    return {m.label: m.value for m in at.metric}


def test_app_modo_demonstracao():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        at = _app({"banco": {"url": f"sqlite:///{Path(tmp, 'vazio.db').as_posix()}"}})
        assert "Ações" not in _metricas(at), "banco vazio: sem painéis"
        at.toggle(key="modo_demo").set_value(True).run()
        assert not at.exception, at.exception
        assert _metricas(at)["Ações"] == "21", _metricas(at)
        assert len(at.tabs) == 8
        at.multiselect[0].select("Centralização").run()
        assert not at.exception, at.exception
        assert _metricas(at)["Ações"] == "7"
        at.toggle(key="modo_demo").set_value(False).run()
        assert "Ações" not in _metricas(at), "modo demonstração não pode gravar no banco do time"


def test_app_le_banco_gravado():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        url = f"sqlite:///{Path(tmp, 'h.db').as_posix()}"
        b = bd.conectar(url)
        _gravar(b, 1, date.today() - timedelta(days=7))
        _gravar(b, 2, date.today())
        b.engine.dispose()
        at = _app({"banco": {"url": url}})
        m = _metricas(at)
        assert m["Ações"] == "21" and "🔁 Prazos replanejados" in m, m
        assert any("2 importação(ões) gravada(s)" in t.value for t in at.markdown)


def test_senha_do_time():
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        at = _app({"app": {"senha": "segredo123"}, "banco": {"url": f"sqlite:///{Path(tmp, 'h.db').as_posix()}"}})
        assert len(at.toggle) == 0, "nada aparece antes da senha"
        at.text_input[0].input("errada")
        at.button[0].click().run()
        assert any("incorreta" in e.value for e in at.error)
        at.text_input[0].input("segredo123")
        at.button[0].click().run()
        assert not at.exception, at.exception
        assert len(at.toggle) > 0, "liberou o app"


TESTES = [test_importador_le_formato_atual, test_sinais_por_prazo, test_avanco_ponderado_igual_planilha,
          test_farol_das_entregas, test_responsaveis_e_qualidade, test_exportacoes_abrem,
          test_planilha_real_se_existir, test_historico_compara_importacoes, test_estado_atual_e_eventos,
          test_recusar_renomeacao_e_voltar, test_desfazer_ultima_importacao, test_banco_em_arquivo_persiste,
          test_app_modo_demonstracao, test_app_le_banco_gravado, test_senha_do_time]

if __name__ == "__main__":
    ok = 0
    for t in TESTES:
        try:
            t()
            ok += 1
            print(f"OK    {t.__name__}")
        except Exception as e:  # noqa: BLE001
            print(f"FALHA {t.__name__}: {e!r}")
    print(f"\n{ok}/{len(TESTES)} testes passaram")
    sys.exit(0 if ok == len(TESTES) else 1)
