"""Leitura da planilha de demandas no formato atual da equipe (Fase 1).

Formato esperado (o mesmo da "Gestão 4.0 - NTec.xlsx"):
- uma aba por FRENTE (Fábrica de Placas, Centralização, Novos Negócios...);
- uma linha de cabeçalho com "Tarefa" na coluna C;
- linha em negrito COM fundo colorido = ENTREGA (tarefa principal);
- linha em negrito SEM fundo = SUBGRUPO dentro da entrega;
- demais linhas = AÇÕES, com Responsável (D), Status (E), Peso (F),
  % (G), Data prevista/conclusão (H) e comentários semanais (I em diante).

As fórmulas da planilha (avanço, data da entrega) NÃO são usadas: o app
recalcula tudo a partir das ações, porque várias linhas-mãe trazem #NUM! ou
#DIV/0!.
"""
from __future__ import annotations

import io
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl
import pandas as pd

COL_AVANCO, COL_TAREFA, COL_RESP, COL_STATUS, COL_PESO, COL_PCT, COL_DATA = 2, 3, 4, 5, 6, 7, 8
COL_SEMANAS = 9
LINHAS_BUSCA_CABECALHO = 20

# Linhas de modelo que vêm vazias na planilha ("AÇÕES" sem responsável etc.)
NOMES_MODELO = {"AÇÕES", "TAREFA PRINCIPAL"}
STATUS_VALIDOS = ("À iniciar", "Em andamento", "Concluído")
SEM_STATUS = "Sem status"
SEM_ENTREGA = "(sem entrega)"

COLUNAS = ["frente", "entrega", "subgrupo", "acao", "responsavel", "status", "peso", "pct",
           "data", "comentario", "historico", "nota", "aba", "linha"]


@dataclass
class Resultado:
    acoes: pd.DataFrame
    abas_lidas: list[str] = field(default_factory=list)
    abas_vazias: list[str] = field(default_factory=list)   # só têm o modelo, sem ação real
    abas_ignoradas: list[str] = field(default_factory=list)  # sem cabeçalho "Tarefa"
    avisos: list[str] = field(default_factory=list)          # problemas de estrutura


def _texto(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _data(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for formato in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip(), formato).date()
            except ValueError:
                pass
    return None


def _tem_fundo(celula) -> bool:
    preench = celula.fill
    if preench is None or not preench.fill_type:
        return False
    cor = preench.fgColor
    if cor is None:
        return False
    if cor.type == "rgb":
        return cor.rgb not in (None, "00000000", "FFFFFFFF")
    return True  # cor de tema/indexada: considera preenchido


def _normalizar_status(v) -> str:
    t = _texto(v)
    if not t:
        return SEM_STATUS
    for s in STATUS_VALIDOS:
        if t.lower() == s.lower() or t.lower().replace("a iniciar", "à iniciar") == s.lower():
            return s
    return t


def _pct(status: str) -> float:
    return {"Concluído": 1.0, "Em andamento": 0.5}.get(status, 0.0)


def _linha_cabecalho(ws) -> int | None:
    for r in range(1, min(LINHAS_BUSCA_CABECALHO, ws.max_row) + 1):
        if _texto(ws.cell(r, COL_TAREFA).value).lower() == "tarefa":
            return r
    return None


def _ler_aba(ws, res: Resultado) -> list[dict]:
    cab = _linha_cabecalho(ws)
    if cab is None:
        res.abas_ignoradas.append(ws.title)
        return []
    semanas = []
    c = COL_SEMANAS
    while c <= ws.max_column and ws.cell(cab, c).value not in (None, ""):
        semanas.append((c, _texto(ws.cell(cab, c).value)))
        c += 1

    linhas, avisos = [], []
    entrega, subgrupo = None, None
    for r in range(cab + 1, ws.max_row + 1):
        celula = ws.cell(r, COL_TAREFA)
        nome = _texto(celula.value)
        if not nome:
            continue
        negrito = bool(celula.font and celula.font.bold)
        if negrito and _tem_fundo(celula):
            entrega, subgrupo = nome, None
            if nome.upper() == "TAREFA PRINCIPAL":
                avisos.append(f"{ws.title}: entrega sem nome (\"TAREFA PRINCIPAL\") na linha {r}")
            d = _data(ws.cell(r, COL_DATA).value)
            if d and d.year < 2025:
                avisos.append(f"{ws.title} › {nome}: data da entrega {d:%d/%m/%Y} parece digitada errada (linha {r})")
            continue
        if negrito:
            subgrupo = nome
            continue

        resp = _texto(ws.cell(r, COL_RESP).value)
        if nome.upper() in NOMES_MODELO and not resp:
            continue
        status = _normalizar_status(ws.cell(r, COL_STATUS).value)
        peso = ws.cell(r, COL_PESO).value
        peso = float(peso) if isinstance(peso, (int, float)) and not isinstance(peso, bool) else None
        historico = [(sem, _texto(ws.cell(r, col).value)) for col, sem in semanas if _texto(ws.cell(r, col).value)]
        nota = ws.cell(r, COL_AVANCO).value
        linhas.append({
            "frente": ws.title.strip(),
            "entrega": entrega or SEM_ENTREGA,
            "subgrupo": subgrupo or "",
            "acao": nome,
            "responsavel": resp,
            "status": status,
            "peso": peso,
            "pct": _pct(status),
            "data": _data(ws.cell(r, COL_DATA).value),
            "comentario": historico[-1][1] if historico else "",
            "historico": historico,
            "nota": nota.strip() if isinstance(nota, str) and not nota.startswith("#") else "",
            "aba": ws.title,
            "linha": r,
        })
    if linhas:
        res.abas_lidas.append(ws.title)
        res.avisos.extend(avisos)  # aba só com o modelo não gera aviso
    else:
        res.abas_vazias.append(ws.title)
    return linhas


def ler_planilha(arquivo) -> Resultado:
    """Lê a planilha (caminho, bytes ou arquivo aberto) e devolve uma ação por linha."""
    if isinstance(arquivo, (bytes, bytearray)):
        arquivo = io.BytesIO(arquivo)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # extensões de formatação condicional não suportadas
        wb = openpyxl.load_workbook(arquivo, data_only=True)
    res = Resultado(acoes=pd.DataFrame(columns=COLUNAS))
    linhas = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        linhas.extend(_ler_aba(ws, res))
    if linhas:
        df = pd.DataFrame(linhas, columns=COLUNAS)
        df.insert(0, "id", range(1, len(df) + 1))
        res.acoes = df
    return res
